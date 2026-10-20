#!/usr/bin/env python3
"""Convierte un archivo Markdown estructurado en un documento Word (.docx) con estilo."""

import re
import sys
from pathlib import Path
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


def set_cell_background(cell, fill_hex: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill_hex)
    tc_pr.append(shd)


LATEX_REPLACEMENTS = [
    (r"\\mathbb\{C\}", "ℂ"),
    (r"\\mathbb\{R\}", "ℝ"),
    (r"\\mathbb\{Z\}", "ℤ"),
    (r"\\mathbb\{N\}_0", "ℕ₀"),
    (r"\\mathbb\{N\}", "ℕ"),
    (r"\\Delta", "Δ"),
    (r"\\alpha", "α"),
    (r"\\beta", "β"),
    (r"\\infty", "∞"),
    (r"\\emptyset", "∅"),
    (r"\\setminus", "∖"),
    (r"\\le\b", "≤"),
    (r"\\ge\b", "≥"),
    (r"\\neq\b", "≠"),
    (r"\\pm\b", "±"),
    (r"\\cdot\b", "·"),
    (r"\\times\b", "×"),
    (r"\\in\b", "∈"),
    (r"\\cup\b", "∪"),
    (r"\\cap\b", "∩"),
    (r"\\iff\b", "⟺"),
    (r"\\implies\b", "⟹"),
    (r"\\bar\{z\}", "z̄"),
    (r"\\bar\{w\}", "w̄"),
    (r"\\left", ""),
    (r"\\right", ""),
    (r"\\log", "log"),
    (r"\\\{", "{"),
    (r"\\\}", "}"),
    (r"\\text\{([^}]*)\}", r"\1"),
    (r"\^\{([^{}]*)\}", r"^\1"),
    (r"_\{([^{}]*)\}", r"_\1"),
    (r"\\sqrt\{([^{}]*)\}", r"√(\1)"),
    (r"\\frac\{([^{}]*)\}\{([^{}]*)\}", r"(\1)/(\2)"),
    (r"\\frac\{([^{}]*)\}\{([^{}]*)\}", r"(\1)/(\2)"),
    (r"\\sqrt\{([^{}]*)\}", r"√(\1)"),
]


def clean_math_text(text: str) -> str:
    for pattern, repl in LATEX_REPLACEMENTS:
        text = re.sub(pattern, repl, text)
    text = text.replace("$", "")
    return text


def add_formatted_runs(paragraph, text: str) -> None:
    """Procesa negritas (**texto**), cursivas (*texto*) y código/fórmula inline (`texto`)."""
    token_re = re.compile(r"(\*\*.*?\*\*|\*[^*]+?\*|`[^`]+?`)")
    pos = 0
    for match in token_re.finditer(text):
        if match.start() > pos:
            paragraph.add_run(clean_math_text(text[pos : match.start()]))
        token = match.group(0)
        if token.startswith("**") and token.endswith("**"):
            run = paragraph.add_run(clean_math_text(token[2:-2]))
            run.bold = True
        elif token.startswith("*") and token.endswith("*"):
            run = paragraph.add_run(clean_math_text(token[1:-1]))
            run.italic = True
        elif token.startswith("`") and token.endswith("`"):
            run = paragraph.add_run(clean_math_text(token[1:-1]))
            run.font.name = "Cambria Math"
            run.font.size = Pt(10.5)
            run.font.color.rgb = RGBColor.from_string("1F4E79")
        pos = match.end()
    if pos < len(text):
        paragraph.add_run(clean_math_text(text[pos:]))


def configure_styles(doc: Document) -> None:
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.15

    for level, (size, color_hex) in enumerate(
        [(16, "1F4E79"), (13, "2E75B6"), (11.5, "1F4E79")], start=1
    ):
        style = doc.styles[f"Heading {level}"]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color_hex)
        style.paragraph_format.space_before = Pt(12 if level < 3 else 8)
        style.paragraph_format.space_after = Pt(4)


def flush_table(doc: Document, table_rows: list[list[str]]) -> None:
    if not table_rows:
        return
    cols = max(len(r) for r in table_rows)
    table = doc.add_table(rows=len(table_rows), cols=cols)
    table.style = "Table Grid"
    for r_idx, row_data in enumerate(table_rows):
        for c_idx in range(cols):
            cell = table.cell(r_idx, c_idx)
            val = row_data[c_idx] if c_idx < len(row_data) else ""
            cell.text = ""
            p = cell.paragraphs[0]
            add_formatted_runs(p, val.strip())
            if r_idx == 0:
                set_cell_background(cell, "1F4E79")
                for run in p.runs:
                    run.bold = True
                    run.font.color.rgb = RGBColor(255, 255, 255)
    doc.add_paragraph()


def convert_markdown_to_docx(md_path: Path, docx_path: Path) -> None:
    doc = Document()
    configure_styles(doc)

    lines = md_path.read_text(encoding="utf-8").splitlines()
    in_code_block = False
    code_lines: list[str] = []
    table_rows: list[list[str]] = []
    title_added = False

    for raw_line in lines:
        line = raw_line.rstrip("\r")

        if line.strip().startswith("```"):
            if table_rows:
                flush_table(doc, table_rows)
                table_rows = []
            if in_code_block:
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Inches(0.3)
                run = p.add_run(clean_math_text("\n".join(code_lines)))
                run.font.name = "Consolas"
                run.font.size = Pt(9.5)
                code_lines = []
                in_code_block = False
            else:
                in_code_block = True
            continue

        if in_code_block:
            code_lines.append(line)
            continue

        # Tablas Markdown
        if line.strip().startswith("|") and line.strip().endswith("|"):
            cells = [c.strip() for c in line.strip()[1:-1].split("|")]
            if all(re.match(r"^:?-{2,}:?$", c) for c in cells if c):
                continue
            table_rows.append(cells)
            continue
        elif table_rows:
            flush_table(doc, table_rows)
            table_rows = []

        stripped = line.strip()
        if not stripped:
            continue

        if stripped in ("---", "***", "___"):
            continue

        if stripped.startswith("# ") and not title_added:
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(12)
            run = p.add_run(stripped[2:].strip())
            run.font.name = "Calibri"
            run.font.size = Pt(18)
            run.bold = True
            run.font.color.rgb = RGBColor.from_string("1F4E79")
            title_added = True
        elif stripped.startswith("# "):
            p = doc.add_heading(level=1)
            add_formatted_runs(p, stripped[2:].strip())
        elif stripped.startswith("## "):
            p = doc.add_heading(level=1)
            add_formatted_runs(p, stripped[3:].strip())
        elif stripped.startswith("### "):
            p = doc.add_heading(level=2)
            add_formatted_runs(p, stripped[4:].strip())
        elif stripped.startswith("#### "):
            p = doc.add_heading(level=3)
            add_formatted_runs(p, stripped[5:].strip())
        elif re.match(r"^[-*+]\s+", stripped):
            p = doc.add_paragraph(style="List Bullet")
            add_formatted_runs(p, re.sub(r"^[-*+]\s+", "", stripped))
        elif re.match(r"^\d+[\.\)]\s+", stripped):
            p = doc.add_paragraph(style="List Number")
            add_formatted_runs(p, re.sub(r"^\d+[\.\)]\s+", "", stripped))
        elif stripped.startswith("> "):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.35)
            add_formatted_runs(p, stripped[2:].strip())
            for run in p.runs:
                run.italic = True
        else:
            p = doc.add_paragraph()
            add_formatted_runs(p, stripped)

    if table_rows:
        flush_table(doc, table_rows)

    docx_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(docx_path))


if __name__ == "__main__":
    if len(sys.argv) != 3:
        print("Uso: python md_to_docx.py <entrada.md> <salida.docx>")
        sys.exit(1)
    convert_markdown_to_docx(Path(sys.argv[1]), Path(sys.argv[2]))
    print(f"Documento generado exitosamente: {sys.argv[2]}")
